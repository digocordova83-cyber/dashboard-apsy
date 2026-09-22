from pathlib import Path
import json
from openpyxl import load_workbook

path = Path('/home/ubuntu/entregas/Dashboard_APSY_Exportacao_Completa.xlsx')
if not path.exists():
    raise SystemExit(f'Arquivo não encontrado: {path}')

wb = load_workbook(path, read_only=True, data_only=True)
result = {
    'file': str(path),
    'size_bytes': path.stat().st_size,
    'sheets': [],
}
for ws in wb.worksheets:
    headers = []
    if ws.max_row >= 1:
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    result['sheets'].append({
        'title': ws.title,
        'rows': ws.max_row,
        'columns': ws.max_column,
        'headers': headers,
    })

out = Path('/home/ubuntu/entregas/auditoria_snapshot_relatorio_apsy.json')
out.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(result, ensure_ascii=False, indent=2))
