from pathlib import Path
from collections import defaultdict
import json
from openpyxl import load_workbook

PATH = Path('/home/ubuntu/entregas/Dashboard_APSY_Exportacao_Completa.xlsx')
OUT = Path('/home/ubuntu/entregas/apsy_media_fronts_snapshot.json')
wb = load_workbook(PATH, read_only=True, data_only=True)

def norm(v):
    return '' if v is None else str(v).strip()

def num(v):
    if v is None or v == '':
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0

def find_header(ws, required):
    for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        vals = [norm(v) for v in row]
        if all(name in vals for name in required):
            return idx, {name: vals.index(name) for name in vals if name}
    raise ValueError(f'Cabeçalho não encontrado: {ws.title} / {required}')

def aggregate(ws_name, required, group_field, metrics):
    ws = wb[ws_name]
    header_row, idx = find_header(ws, required)
    out = defaultdict(lambda: defaultdict(float))
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        group = norm(row[idx[group_field]]) or '(vazio)'
        for metric in metrics:
            out[group][metric] += num(row[idx[metric]]) if metric in idx else 0.0
        out[group]['linhas'] += 1
    return {k: dict(v) for k, v in out.items()}

meta_by_objective = aggregate(
    'Meta_Campanhas',
    ['Campaign', 'Objetivo da campanha', 'Spend'],
    'Objetivo da campanha',
    ['Spend', 'Clicks', 'Impressions', 'Reach', 'Actions Lead', 'Actions Offsite Conversion Fb Pixel Lead', 'Actions Onsite Conversion Messaging Conversation Started 7D'],
)
google_by_type = aggregate(
    'Google_Campanhas',
    ['Tipo de canal', 'Campaign', 'Spend'],
    'Tipo de canal',
    ['Spend', 'Clicks', 'Impressions', 'Conversions'],
)

# Totais de programática a partir das abas diárias persistidas.
prog = {}
for sheet, label in [('Prog_DV360', 'DV360'), ('Prog_Push', 'Push')]:
    ws = wb[sheet]
    header_row = None
    idx = None
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        vals = [norm(v) for v in row]
        if 'Spend' in vals:
            header_row = i
            idx = {name: vals.index(name) for name in vals if name}
            break
    if header_row is None:
        raise ValueError(f'Cabeçalho não encontrado em {sheet}')
    totals = defaultdict(float)
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        for metric in ['Spend', 'Impressions', 'Clicks', 'Dispatches', 'Complete Views']:
            if metric in idx:
                totals[metric] += num(row[idx[metric]])
    prog[label] = dict(totals)

result = {
    'snapshot': str(PATH),
    'meta_by_objective': meta_by_objective,
    'google_by_channel_type': google_by_type,
    'programmatic': prog,
}
OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(result, ensure_ascii=False, indent=2))
