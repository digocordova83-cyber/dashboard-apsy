from pathlib import Path
from collections import Counter, defaultdict
import json
from openpyxl import load_workbook

PATH = Path('/home/ubuntu/entregas/Dashboard_APSY_Exportacao_Completa.xlsx')
OUT = Path('/home/ubuntu/entregas/apsy_snapshot_efficiency_audit.json')

wb = load_workbook(PATH, read_only=True, data_only=True)

def rows_values(ws):
    for row in ws.iter_rows(values_only=True):
        yield list(row)

def find_header(ws, required):
    for idx, row in enumerate(rows_values(ws), start=1):
        normalized = {str(v).strip() for v in row if v is not None}
        if all(name in normalized for name in required):
            return idx, row
    raise ValueError(f'Cabeçalho não encontrado em {ws.title}: {required}')

def norm(value):
    return '' if value is None else str(value).strip()

def classify_utm(value):
    v = norm(value).lower()
    if v in {'fb', 'ig', 'meta', 'meta-ads', 'facebook', 'instagram'}:
        return 'Meta explícito'
    if v in {'google', 'google-ads', 'google ads'}:
        return 'Google explícito'
    return 'Outras/vazias'

crm = wb['CRM_Leads']
header_row, headers = find_header(crm, ['Oportunidade (nº)', 'Oportunidade (etapa)', 'Canal de origem', 'UTM source'])
index = {norm(v): i for i, v in enumerate(headers) if norm(v)}
records = {}
for row in crm.iter_rows(min_row=header_row + 1, values_only=True):
    opp = norm(row[index['Oportunidade (nº)']])
    stage = norm(row[index['Oportunidade (etapa)']]).upper()
    if not opp or not stage or stage == 'OUTROS':
        continue
    records[opp] = {
        'stage': stage,
        'source_channel': norm(row[index['Canal de origem']]) or 'desconhecido',
        'utm_source': norm(row[index['UTM source']]),
        'utm_medium': norm(row[index['UTM medium']]),
        'utm_campaign': norm(row[index['UTM campaign']]),
    }

by_source = defaultdict(lambda: Counter())
by_utm = defaultdict(lambda: Counter())
coverage = Counter()
for rec in records.values():
    for bucket in (by_source[rec['source_channel']], by_utm[classify_utm(rec['utm_source'])]):
        bucket['oportunidades'] += 1
        if rec['stage'] == 'SQL':
            bucket['sqls'] += 1
        if rec['stage'] == 'MATRICULADO':
            bucket['matriculados_crm'] += 1
    coverage['oportunidades'] += 1
    if rec['utm_source']:
        coverage['oportunidades_com_utm_source'] += 1
    if rec['utm_campaign']:
        coverage['oportunidades_com_utm_campaign'] += 1
    if rec['stage'] == 'SQL':
        coverage['sqls'] += 1
        if rec['utm_source']:
            coverage['sqls_com_utm_source'] += 1
    if rec['stage'] == 'MATRICULADO':
        coverage['matriculados_crm'] += 1
        if rec['utm_source']:
            coverage['matriculados_crm_com_utm_source'] += 1

sheet_samples = {}
for title in ['Meta_Diario', 'Meta_Campanhas', 'Google_Diario', 'Google_Campanhas', 'Prog_DV360', 'Prog_Push']:
    ws = wb[title]
    sample = []
    for i, row in enumerate(rows_values(ws), start=1):
        if any(v is not None for v in row):
            sample.append({'row': i, 'values': row[:20]})
        if len(sample) >= 8:
            break
    sheet_samples[title] = sample

result = {
    'snapshot': str(PATH),
    'crm_header_row': header_row,
    'crm_valid_opportunities': len(records),
    'coverage': dict(coverage),
    'by_source_channel': {k: dict(v) for k, v in sorted(by_source.items(), key=lambda kv: -kv[1]['oportunidades'])},
    'by_utm_group': {k: dict(v) for k, v in sorted(by_utm.items(), key=lambda kv: -kv[1]['oportunidades'])},
    'sheet_samples': sheet_samples,
}
OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
