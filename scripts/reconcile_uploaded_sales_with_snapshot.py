from pathlib import Path
from collections import Counter, defaultdict
import json
import re
import unicodedata
from openpyxl import load_workbook

SALES_PATH = Path('/home/ubuntu/upload/pasted_file_C21JVL_DashboardPósGraduação.xlsx')
CRM_PATH = Path('/home/ubuntu/entregas/Dashboard_APSY_Exportacao_Completa.xlsx')
OUT = Path('/home/ubuntu/entregas/apsy_sales_crm_reconciliation_20260918.json')

def norm(value):
    text = '' if value is None else str(value)
    text = ''.join(ch for ch in unicodedata.normalize('NFD', text) if unicodedata.category(ch) != 'Mn')
    text = re.sub(r'[^a-zA-Z0-9]+', ' ', text.lower()).strip()
    return re.sub(r'\s+', ' ', text)

def course_tokens(course):
    text = norm(course)
    if 'neuropsic' in text:
        return ['neuropsic', 'avaliacao neuro']
    if 'modelos contemporaneos' in text or 'tcc' in text:
        return ['tcc', 'modelos contemporaneos', 'cognitivo comportamental']
    if 'baseada em processos' in text:
        return ['tbp', 'baseada em processos']
    return [token for token in text.split() if len(token) >= 4]

def source_group(rec):
    text = norm(' '.join([rec.get('Canal de origem',''), rec.get('UTM source',''), rec.get('UTM medium',''), rec.get('UTM campaign','')]))
    if re.search(r'google|cpc|search|gads', text):
        return 'Google'
    if re.search(r'facebook|instagram|meta|whatsapp|fb|ig', text):
        return 'Meta / WhatsApp'
    if re.search(r'organico|organic|direto|direct', text):
        return 'Orgânico / Direto'
    return 'Outros / não identificado'

def candidate_score(rec, course):
    context = norm(' '.join([rec.get('Products',''), rec.get('Oportunidade (nome)',''), rec.get('Oportunidade (tag)',''), rec.get('UTM campaign',''), rec.get('Formulário',''), rec.get('Contexto adicional','')]))
    score = 100 if any(token in context for token in course_tokens(course)) else 0
    stage = norm(rec.get('Oportunidade (etapa)','')).upper()
    score += {'MATRICULADO':40, 'SQL':30, 'SAL':20, 'MQL':10}.get(stage, 0)
    if stage and stage != 'OUTROS':
        score += 5
    return score

sales_wb = load_workbook(SALES_PATH, read_only=True, data_only=True)
ws = sales_wb['Vendas']
rows = ws.iter_rows(values_only=True)
headers = [str(v).strip() if v is not None else '' for v in next(rows)]
idx = {name: i for i, name in enumerate(headers)}
sales = []
for row in rows:
    if str(row[idx['Status']]).strip() != 'Ativa':
        continue
    sales.append({
        'code': str(row[idx['Código']]).strip(),
        'name': str(row[idx['Cliente']]).strip(),
        'course': str(row[idx['Serviço']]).strip(),
        'gross': float(row[idx['Faturamento Bruto']] or 0),
        'net': float(row[idx['Faturamento Líquido']] or 0),
    })

crm_wb = load_workbook(CRM_PATH, read_only=True, data_only=True)
crm_ws = crm_wb['CRM_Leads']
crm_headers = None
header_row = None
for i, row in enumerate(crm_ws.iter_rows(values_only=True), start=1):
    vals = [str(v).strip() if v is not None else '' for v in row]
    if 'Contato' in vals and 'Oportunidade (etapa)' in vals:
        crm_headers = vals
        header_row = i
        break
crm_idx = {name: i for i, name in enumerate(crm_headers)}
crm_records = []
for row in crm_ws.iter_rows(min_row=header_row+1, values_only=True):
    rec = {name: ('' if row[i] is None else str(row[i]).strip()) for name, i in crm_idx.items()}
    if rec.get('Contato'):
        crm_records.append(rec)

by_name = defaultdict(list)
for rec in crm_records:
    by_name[norm(rec['Contato'])].append(rec)

reconciled = []
for sale in sales:
    candidates = by_name.get(norm(sale['name']), [])
    ranked = sorted(candidates, key=lambda rec: (candidate_score(rec, sale['course']), rec.get('Atualizado em','')), reverse=True)
    chosen = ranked[0] if ranked else None
    reconciled.append({
        **sale,
        'match_count': len(candidates),
        'matched': bool(chosen),
        'crm': None if not chosen else {
            'source_channel': chosen.get('Canal de origem'),
            'utm_source': chosen.get('UTM source'),
            'utm_campaign': chosen.get('UTM campaign'),
            'stage': chosen.get('Oportunidade (etapa)'),
            'source_group': source_group(chosen),
        },
    })

matched = [row for row in reconciled if row['matched']]
count = lambda values: dict(Counter(values))
by_course = {}
for course in sorted(set(row['course'] for row in sales)):
    subset = [row for row in reconciled if row['course'] == course]
    subset_matched = [row for row in subset if row['matched']]
    by_course[course] = {
        'active_sales': len(subset),
        'matched_in_crm': len(subset_matched),
        'coverage_pct': round(100*len(subset_matched)/len(subset),1),
    }

result = {
    'basis': 'Vendas com status Ativa na nova base; correspondência exata por nome completo normalizado contra o snapshot CRM de 13/09/2026.',
    'active_sales': len(sales),
    'matched_in_crm': len(matched),
    'unmatched': len(sales)-len(matched),
    'coverage_pct': round(100*len(matched)/len(sales),1),
    'multiple_name_matches': sum(1 for row in reconciled if row['match_count'] > 1),
    'matched_with_utm_source': sum(1 for row in matched if row['crm']['utm_source']),
    'by_course': by_course,
    'matched_by_source_group': count(row['crm']['source_group'] for row in matched),
    'matched_by_source_channel': count(row['crm']['source_channel'] for row in matched),
    'matched_by_stage': count(row['crm']['stage'] for row in matched),
    'unmatched_sales': [{k: row[k] for k in ['code','name','course']} for row in reconciled if not row['matched']],
}
OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(result, ensure_ascii=False, indent=2))
