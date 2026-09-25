import json
import os
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SOURCE_FILE = Path("/home/ubuntu/entregas/Dashboard_APSY_Exportacao_Completa_Dados.json")
OUTPUT_FILE = Path("/home/ubuntu/entregas/Dashboard_APSY_Exportacao_Completa.xlsx")

THEME = {
    "primary": "013C4C",
    "secondary": "00ACB3",
    "accent": "F26457",
    "paper": "FFFFFF",
    "mist": "EFF6F7",
    "line": "D7E3E6",
    "text": "172B32",
    "muted": "61737A",
    "warning": "FFF5E1",
    "success": "E9F7F4",
}

TABLE_DESCRIPTIONS = {
    "CRM_Leads": "Base completa do EducaCRM persistida no Dashboard APSY. O funil usa a taxonomia oficial única; testes e atendimento em outro canal ficam em FORA_DA_BASE.",
    "Meta_Diario": "Dados diários de Meta Ads por campanha. Leads replicam a regra do dashboard, sem dupla contagem de ações de lead e sem tratar conversas de WhatsApp como leads.",
    "Meta_Campanhas": "Performance consolidada por campanha Meta no período de mídia informado.",
    "Meta_Conjuntos": "Performance consolidada por conjunto de anúncios Meta no período de mídia informado.",
    "Meta_Anuncios": "Performance consolidada por anúncio Meta, incluindo thumbnail ou URL de mídia quando disponível.",
    "Meta_Demografia": "Dados de faixa etária e gênero retornados pela integração Meta/Windsor.",
    "Google_Diario": "Dados diários de Google Ads por campanha.",
    "Google_Campanhas": "Performance consolidada por campanha Google Ads, com dados de parcela de impressões quando disponíveis.",
    "Google_Palavras": "Dados brutos de palavras-chave retornados pela integração Google Ads/Windsor. Uma palavra pode ter mais de uma linha de origem.",
    "Google_Termos": "Dados brutos de termos de pesquisa retornados pela integração Google Ads/Windsor.",
    "Google_Dispositivos": "Dados de Google Ads por dispositivo e campanha.",
    "Prog_DV360": "Snapshot persistido de dados programáticos DV360 por insertion order e dia.",
    "Prog_Meta_Social": "Snapshot persistido de social Meta no bloco programático por campanha e dia.",
    "Prog_Push": "Snapshot persistido de Push Notification por dia.",
    "Publya_Campanhas": "Campanhas retornadas pela API Publya no momento da exportação. Métricas programáticas detalhadas estão nas abas de snapshot persistido.",
    "Metas": "Metas configuradas dentro do Dashboard APSY.",
    "Otimizacoes": "Diagnósticos, defesas e planos de ação registrados no módulo de otimizações.",
    "Auditoria": "Log de auditoria das ações registradas no dashboard.",
    "Cache_Manifest": "Metadados de cache Windsor: chave, horário de busca e tamanho aproximado do payload. O conteúdo de mídia está exportado em abas próprias para evitar duplicação.",
}

DISPLAY_NAMES = {
    "CRM_Leads": "CRM — Leads completos",
    "Meta_Diario": "Meta — Diário",
    "Meta_Campanhas": "Meta — Campanhas",
    "Meta_Conjuntos": "Meta — Conjuntos",
    "Meta_Anuncios": "Meta — Anúncios",
    "Meta_Demografia": "Meta — Demografia",
    "Google_Diario": "Google — Diário",
    "Google_Campanhas": "Google — Campanhas",
    "Google_Palavras": "Google — Palavras-chave",
    "Google_Termos": "Google — Termos de busca",
    "Google_Dispositivos": "Google — Dispositivos",
    "Prog_DV360": "Programática — DV360",
    "Prog_Meta_Social": "Programática — Meta Social",
    "Prog_Push": "Programática — Push",
    "Publya_Campanhas": "Programática — Campanhas Publya",
    "Metas": "Configuração — Metas",
    "Otimizacoes": "Otimizações",
    "Auditoria": "Auditoria",
    "Cache_Manifest": "Metadados — Cache Windsor",
}

DATE_COLUMNS = {"date", "day", "createdDate", "updatedDate", "importedAt", "createdAt", "updatedAt", "fetchedAt", "completedAt", "dueDate", "lastLogin"}
PERCENT_COLUMNS = {"frequency", "viewability", "completionRate", "vcr", "search_impression_share", "search_budget_lost_impression_share", "search_rank_lost_impression_share"}
CURRENCY_HINTS = ("spend", "budget", "revenue", "value", "cost", "cpv", "cpc", "cpm", "cpl", "cpa", "cac", "targetValue", "saleValue")


def stringify(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return value


def columns_for(rows: list[dict[str, Any]]) -> list[str]:
    cols: list[str] = []
    for row in rows:
        for key in row.keys():
            if key not in cols:
                cols.append(key)
    return cols


def friendly_header(name: str) -> str:
    replacements = {
        "externalId": "ID externo", "companyName": "Empresa", "contactName": "Contato", "sourceChannel": "Canal de origem",
        "formName": "Formulário", "utmSource": "UTM source", "utmMedium": "UTM medium", "utmCampaign": "UTM campaign",
        "addrNumber": "Número", "birthDate": "Data de nascimento", "extraContext": "Contexto adicional",
        "opportunityNumber": "Oportunidade (nº)", "opportunityName": "Oportunidade (nome)", "opportunityTag": "Oportunidade (tag)",
        "opportunityStage": "Oportunidade (etapa)", "createdDate": "Criado em", "updatedDate": "Atualizado em", "importedAt": "Importado em",
        "adset_name": "Conjunto", "ad_group_name": "Grupo de anúncios", "ad_name": "Anúncio", "search_term": "Termo de pesquisa",
        "keyword_text": "Palavra-chave", "campaign_status": "Status da campanha", "campaign_objective": "Objetivo da campanha",
        "advertising_channel_type": "Tipo de canal", "conversion_value": "Valor de conversão", "cacheKey": "Chave de cache",
        "fetchedAt": "Atualizado em", "payloadRows": "Linhas no payload", "payloadType": "Tipo de payload",
    }
    return replacements.get(name, name.replace("_", " ").replace("-", " ").strip().title())


def value_number_format(col_name: str, value: Any) -> str | None:
    n = col_name.lower()
    if col_name in PERCENT_COLUMNS or "share" in n or "rate" in n or "viewability" in n:
        return "0.00%"
    if any(token.lower() in n for token in CURRENCY_HINTS):
        return 'R$ #,##0.00'
    if isinstance(value, float):
        return '#,##0.00'
    if isinstance(value, int):
        return '#,##0'
    return None


def create_data_sheet(wb: Workbook, sheet_name: str, rows: list[dict[str, Any]], description: str) -> None:
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=8)
    ws["B2"] = DISPLAY_NAMES.get(sheet_name, sheet_name)
    ws["B2"].font = Font(name="Georgia", size=18, bold=True, color=THEME["primary"])
    ws["B2"].alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 30
    ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=20)
    ws["B3"] = description
    ws["B3"].font = Font(name="Calibri", size=10, italic=True, color=THEME["muted"])
    ws["B3"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[3].height = 32

    if not rows:
        ws["B5"] = "Sem registros disponíveis nesta fonte no momento da exportação."
        ws["B5"].font = Font(name="Calibri", size=11, italic=True, color=THEME["muted"])
        return

    columns = columns_for(rows)
    header_row = 5
    for c_index, col_name in enumerate(columns, start=2):
        cell = ws.cell(row=header_row, column=c_index, value=friendly_header(col_name))
        cell.font = Font(name="Georgia", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=THEME["primary"])
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 30

    thin = Side(style="thin", color=THEME["line"])
    for r_index, row in enumerate(rows, start=header_row + 1):
        for c_index, col_name in enumerate(columns, start=2):
            raw_value = stringify(row.get(col_name))
            cell = ws.cell(row=r_index, column=c_index, value=raw_value)
            cell.font = Font(name="Calibri", size=10, color=THEME["text"])
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
            if isinstance(raw_value, (int, float)) and not isinstance(raw_value, bool):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = value_number_format(col_name, raw_value) or "#,##0.00"
            elif col_name in DATE_COLUMNS and raw_value:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=thin)
            if r_index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F8FBFC")
        ws.row_dimensions[r_index].height = 18

    last_col = get_column_letter(len(columns) + 1)
    last_row = header_row + len(rows)
    ws.auto_filter.ref = f"B{header_row}:{last_col}{last_row}"
    ws.freeze_panes = f"B{header_row + 1}"

    for c_index, col_name in enumerate(columns, start=2):
        values = [str(stringify(row.get(col_name)) or "") for row in rows]
        max_len = max([len(friendly_header(col_name))] + [min(len(value), 50) for value in values])
        if any(token in col_name.lower() for token in ["notes", "details", "diagnosis", "defense", "actionplan", "checklist", "thumbnail", "image", "url"]):
            width = 42
        elif col_name in DATE_COLUMNS:
            width = 20
        elif any(token in col_name.lower() for token in ["spend", "impression", "click", "conversion", "view", "result", "reach", "frequency", "rate", "score", "value"]):
            width = 16
        else:
            width = min(max(max_len + 3, 13), 42)
        ws.column_dimensions[get_column_letter(c_index)].width = width

    numeric_cols = [i + 2 for i, c in enumerate(columns) if any(t in c.lower() for t in ["spend", "impressions", "clicks", "conversions", "results", "reach"])]
    for col_index in numeric_cols[:3]:
        ws.conditional_formatting.add(
            f"{get_column_letter(col_index)}{header_row + 1}:{get_column_letter(col_index)}{last_row}",
            DataBarRule(start_type="min", end_type="max", color=THEME["secondary"]),
        )


def create_index_sheet(wb: Workbook, data: dict[str, Any]) -> None:
    ws = wb.active
    ws.title = "Índice"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    for col in range(2, 7):
        ws.column_dimensions[get_column_letter(col)].width = [30, 22, 18, 18, 36][col - 2]

    ws.merge_cells("B2:F2")
    ws["B2"] = "Dashboard APSY — Exportação completa"
    ws["B2"].font = Font(name="Georgia", size=22, bold=True, color=THEME["primary"])
    ws["B2"].alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 36
    ws.merge_cells("B3:F3")
    meta = data["metadata"]
    ws["B3"] = f"Gerado em {meta['generatedAtBrt']} (BRT) · Mídia: {meta['mediaDateFrom']} a {meta['mediaDateTo']} · CRM: {meta['crmDateFrom']} a {meta['crmDateTo']}"
    ws["B3"].font = Font(name="Calibri", size=10, color=THEME["muted"])
    ws.row_dimensions[3].height = 22

    ws.merge_cells("B5:F5")
    ws["B5"] = "COBERTURA DA EXPORTAÇÃO"
    ws["B5"].font = Font(name="Georgia", size=12, bold=True, color="FFFFFF")
    ws["B5"].fill = PatternFill("solid", fgColor=THEME["primary"])
    ws["B5"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[5].height = 24

    headers = ["Aba", "Fonte", "Registros", "Período", "Status"]
    for idx, header in enumerate(headers, start=2):
        cell = ws.cell(row=6, column=idx, value=header)
        cell.font = Font(name="Georgia", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=THEME["secondary"])
        cell.alignment = Alignment(horizontal="center", vertical="center")

    source_rows = sorted(data["sourceStatus"], key=lambda x: x["sheet"])
    thin = Side(style="thin", color=THEME["line"])
    for r, source in enumerate(source_rows, start=7):
        sheet = source["sheet"]
        values = [DISPLAY_NAMES.get(sheet, sheet), source["source"], source["rows"], f"{source.get('dateFrom') or '—'} a {source.get('dateTo') or '—'}", source["status"].upper()]
        for c, value in enumerate(values, start=2):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = Font(name="Calibri", size=10, color=THEME["text"])
            cell.alignment = Alignment(horizontal="right" if c == 4 else "left", vertical="center")
            cell.border = Border(bottom=thin)
            if c == 3 and sheet in data["tables"]:
                cell.hyperlink = f"#'{sheet}'!B2"
                cell.font = Font(name="Calibri", size=10, color=THEME["primary"], underline="single")
            if c == 6:
                color = THEME["success"] if source["status"] == "ok" else THEME["warning"]
                cell.fill = PatternFill("solid", fgColor=color)
                cell.font = Font(name="Calibri", size=10, bold=True, color=THEME["text"])
        ws.row_dimensions[r].height = 18

    note_row = 7 + len(source_rows) + 2
    ws.merge_cells(start_row=note_row, start_column=2, end_row=note_row, end_column=6)
    ws.cell(note_row, 2, "NOTAS DE SEGURANÇA E COBERTURA")
    ws.cell(note_row, 2).font = Font(name="Georgia", size=12, bold=True, color="FFFFFF")
    ws.cell(note_row, 2).fill = PatternFill("solid", fgColor=THEME["primary"])
    for offset, text in enumerate(meta["exclusions"], start=1):
        ws.merge_cells(start_row=note_row + offset, start_column=2, end_row=note_row + offset, end_column=6)
        cell = ws.cell(note_row + offset, 2, f"• {text}")
        cell.font = Font(name="Calibri", size=10, color=THEME["muted"])
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[note_row + offset].height = 25
    ws.freeze_panes = "B6"


def create_dictionary_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Dicionário")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 90
    ws.merge_cells("B2:C2")
    ws["B2"] = "Dicionário de dados e critérios de leitura"
    ws["B2"].font = Font(name="Georgia", size=18, bold=True, color=THEME["primary"])
    ws.row_dimensions[2].height = 30

    ws["B4"] = "Aba"
    ws["C4"] = "Conteúdo e definição"
    for cell in ws[4][1:3]:
        cell.font = Font(name="Georgia", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=THEME["primary"])
        cell.alignment = Alignment(horizontal="center", vertical="center")
    row = 5
    thin = Side(style="thin", color=THEME["line"])
    for sheet, description in TABLE_DESCRIPTIONS.items():
        ws.cell(row=row, column=2, value=DISPLAY_NAMES.get(sheet, sheet))
        ws.cell(row=row, column=3, value=description)
        for col in (2, 3):
            cell = ws.cell(row=row, column=col)
            cell.font = Font(name="Calibri", size=10, color=THEME["text"])
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = Border(bottom=thin)
        ws.row_dimensions[row].height = 42
        row += 1
    ws.freeze_panes = "B5"


def main() -> None:
    if not SOURCE_FILE.exists():
        raise FileNotFoundError(f"Arquivo de extração não encontrado: {SOURCE_FILE}")
    with SOURCE_FILE.open("r", encoding="utf-8") as fh:
        data = json.load(fh)

    wb = Workbook()
    create_index_sheet(wb, data)
    for sheet_name, rows in data["tables"].items():
        create_data_sheet(wb, sheet_name, rows if isinstance(rows, list) else [], TABLE_DESCRIPTIONS.get(sheet_name, "Dados exportados do Dashboard APSY."))
    create_dictionary_sheet(wb)

    for ws in wb.worksheets:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_margins.left = 0.25
        ws.page_margins.right = 0.25
        ws.page_margins.top = 0.5
        ws.page_margins.bottom = 0.5

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_FILE)

    check = load_workbook(OUTPUT_FILE, read_only=True, data_only=False)
    result = {
        "output": str(OUTPUT_FILE),
        "sheets": check.sheetnames,
        "sheet_count": len(check.sheetnames),
        "crm_rows": check["CRM_Leads"].max_row - 5,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
