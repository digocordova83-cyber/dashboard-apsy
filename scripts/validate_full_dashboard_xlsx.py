import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

SOURCE = Path("/home/ubuntu/entregas/Dashboard_APSY_Exportacao_Completa.xlsx")
PREVIEW = Path("/home/ubuntu/entregas/Dashboard_APSY_Exportacao_Completa_Preview.xlsx")
REPORT = Path("/home/ubuntu/entregas/Dashboard_APSY_Exportacao_Completa_Validacao.json")


def copy_sheet_sample(source_ws, target_wb, name: str, max_rows: int) -> None:
    target_ws = target_wb.create_sheet(name)
    for row in source_ws.iter_rows(min_row=1, max_row=min(source_ws.max_row, max_rows), values_only=True):
        target_ws.append(list(row))
    for col in range(1, min(source_ws.max_column, 14) + 1):
        target_ws.column_dimensions[chr(64 + col)].width = 18
    target_ws.freeze_panes = "B6" if source_ws.max_row >= 6 else "A1"


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {SOURCE}")
    source = load_workbook(SOURCE, read_only=True, data_only=False)
    expected_sheets = [
        "Índice", "CRM_Leads", "Meta_Diario", "Meta_Campanhas", "Meta_Conjuntos",
        "Meta_Anuncios", "Meta_Demografia", "Google_Diario", "Google_Campanhas",
        "Google_Palavras", "Google_Termos", "Google_Dispositivos", "Prog_DV360",
        "Prog_Meta_Social", "Prog_Push", "Publya_Campanhas", "Metas", "Otimizacoes",
        "Auditoria", "Cache_Manifest", "Dicionário",
    ]
    missing = [sheet for sheet in expected_sheets if sheet not in source.sheetnames]
    data_rows = {sheet: max(source[sheet].max_row - 5, 0) for sheet in source.sheetnames if sheet not in {"Índice", "Dicionário"}}
    checks = {
        "arquivo_existe": SOURCE.exists(),
        "abas_esperadas": len(expected_sheets),
        "abas_encontradas": len(source.sheetnames),
        "abas_ausentes": missing,
        "crm_linhas": data_rows.get("CRM_Leads"),
        "google_termos_linhas": data_rows.get("Google_Termos"),
        "crm_conforme_base": data_rows.get("CRM_Leads") == 6072,
        "google_termos_integrais": data_rows.get("Google_Termos") == 49202,
        "dados_por_aba": data_rows,
        "resultado": "aprovado" if not missing and data_rows.get("CRM_Leads") == 6072 and data_rows.get("Google_Termos") == 49202 else "revisar",
    }
    REPORT.write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8")

    preview = Workbook()
    preview.remove(preview.active)
    copy_sheet_sample(source["Índice"], preview, "Índice", 35)
    copy_sheet_sample(source["CRM_Leads"], preview, "CRM_Amostra", 35)
    copy_sheet_sample(source["Meta_Campanhas"], preview, "Meta_Campanhas", 35)
    copy_sheet_sample(source["Google_Termos"], preview, "Google_Termos_Amostra", 35)
    copy_sheet_sample(source["Dicionário"], preview, "Dicionário", 35)
    preview.save(PREVIEW)
    print(json.dumps({"validation": str(REPORT), "preview": str(PREVIEW), **checks}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
