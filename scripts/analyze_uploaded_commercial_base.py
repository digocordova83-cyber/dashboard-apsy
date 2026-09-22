from pathlib import Path
from collections import defaultdict
import json
from openpyxl import load_workbook

PATH = Path('/home/ubuntu/upload/pasted_file_C21JVL_DashboardPósGraduação.xlsx')
OUT = Path('/home/ubuntu/entregas/apsy_commercial_base_audit_20260918.json')
wb = load_workbook(PATH, read_only=True, data_only=True)

def norm(v):
    return '' if v is None else str(v).strip()

def num(v):
    if v is None or v == '':
        return 0.0
    return float(v)

ws = wb['Vendas']
rows = ws.iter_rows(values_only=True)
headers = [norm(v) for v in next(rows)]
idx = {name: i for i, name in enumerate(headers)}

summary = defaultdict(lambda: {'sales': 0, 'gross': 0.0, 'discount': 0.0, 'net': 0.0, 'scholarships_100': 0})
status = defaultdict(lambda: {'records': 0, 'gross': 0.0, 'discount': 0.0, 'net': 0.0})
active_records = []
for row in rows:
    service = norm(row[idx['Serviço']])
    row_status = norm(row[idx['Status']]) or '(vazio)'
    gross = num(row[idx['Faturamento Bruto']])
    discount = num(row[idx['Desconto']])
    net = num(row[idx['Faturamento Líquido']])
    scholarship = norm(row[idx['Bolsa 100%']]).lower() == 'sim'
    status[row_status]['records'] += 1
    status[row_status]['gross'] += gross
    status[row_status]['discount'] += discount
    status[row_status]['net'] += net
    if row_status == 'Ativa':
        summary[service]['sales'] += 1
        summary[service]['gross'] += gross
        summary[service]['discount'] += discount
        summary[service]['net'] += net
        summary[service]['scholarships_100'] += int(scholarship)
        active_records.append({
            'code': norm(row[idx['Código']]),
            'client': norm(row[idx['Cliente']]),
            'service': service,
            'gross': gross,
            'discount': discount,
            'net': net,
            'discount_pct': num(row[idx['% Desconto']]),
            'scholarship_100': scholarship,
        })

totals = {
    'sales': sum(v['sales'] for v in summary.values()),
    'gross': sum(v['gross'] for v in summary.values()),
    'discount': sum(v['discount'] for v in summary.values()),
    'net': sum(v['net'] for v in summary.values()),
    'scholarships_100': sum(v['scholarships_100'] for v in summary.values()),
}
totals['gross_ticket'] = totals['gross'] / totals['sales'] if totals['sales'] else 0
_tot_nonzero = [r for r in active_records if r['net'] > 0]
totals['net_ticket_including_scholarships'] = totals['net'] / totals['sales'] if totals['sales'] else 0
totals['net_ticket_paying_students'] = totals['net'] / len(_tot_nonzero) if _tot_nonzero else 0
totals['avg_discount_pct'] = totals['discount'] / totals['gross'] if totals['gross'] else 0

for values in summary.values():
    values['gross_ticket'] = values['gross'] / values['sales'] if values['sales'] else 0
    values['net_ticket'] = values['net'] / values['sales'] if values['sales'] else 0
    values['avg_discount_pct'] = values['discount'] / values['gross'] if values['gross'] else 0

result = {
    'source_file': str(PATH),
    'active_totals': totals,
    'active_by_service': dict(summary),
    'status_summary': dict(status),
}

for sheet_name, key in [('Mensal', 'monthly_sheet'), ('Serviços', 'services_sheet'), ('Exclusões', 'exclusions_sheet')]:
    sheet = wb[sheet_name]
    result[key] = [
        [value for value in row]
        for row in sheet.iter_rows(values_only=True)
        if any(value is not None for value in row)
    ]

OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
